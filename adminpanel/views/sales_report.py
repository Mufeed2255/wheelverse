import io
from datetime import datetime
from decimal import Decimal

from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.core.paginator import Paginator
from django.db.models import (
    Count,
    DecimalField,
    ExpressionWrapper,
    F,
    IntegerField,
    Max,
    Q,
    Sum,
    Value,
)
from django.db.models.functions import Coalesce, TruncMonth
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.http import urlencode

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from Orders.models import Order, OrderItem
from adminpanel.models import Category


ZERO_MONEY = Decimal("0.00")

# These statuses represent orders that reached delivery.
# Their sales/discount values remain in the report even if the order
# later moves into a return-related status.
SALES_ORDER_STATUSES = [
    "DELIVERED",
    "RETURN_REQUESTED",
    "RETURN_APPROVED",
    "RETURN_REJECTED",
    "RETURNED",
]


def parse_report_date(date_value):
    if not date_value:
        return None

    try:
        return datetime.strptime(date_value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def calculate_percentage_change(current_value, previous_value):
    current_value = Decimal(current_value or 0)
    previous_value = Decimal(previous_value or 0)

    if previous_value == 0:
        return Decimal("100.00") if current_value > 0 else ZERO_MONEY

    return (
        (current_value - previous_value)
        / previous_value
        * Decimal("100")
    )


def _money(value):
    return Decimal(value or 0).quantize(Decimal("0.01"))


def _build_search_filter(search):
    return (
        Q(product_name__icontains=search)
        | Q(variant__product__name__icontains=search)
        | Q(variant__product__sku__icontains=search)
        | Q(order__order_id__icontains=search)
        | Q(order__user__username__icontains=search)
        | Q(order__user__email__icontains=search)
    )


def _get_discount_summary(
    *,
    start_date,
    end_date,
    search,
    category_id,
):
    """
    Return offer, coupon and total discounts.

    For the all-category report, use exact order-level discount fields.

    For a selected category, allocate each order-level discount
    proportionally according to that category's item value inside the order.
    This prevents the category discount summary from incorrectly becoming zero.
    """

    discount_orders = (
        Order.objects
        .filter(
            ordered_at__date__range=[start_date, end_date],
            status__in=SALES_ORDER_STATUSES,
        )
        .prefetch_related(
            "items",
            "items__variant",
            "items__variant__product",
            "items__variant__product__category",
        )
        .select_related("user")
        .distinct()
    )

    if search:
        discount_orders = discount_orders.filter(
            Q(order_id__icontains=search)
            | Q(user__username__icontains=search)
            | Q(user__email__icontains=search)
            | Q(items__product_name__icontains=search)
            | Q(items__variant__product__name__icontains=search)
            | Q(items__variant__product__sku__icontains=search)
        ).distinct()

    if category_id == "all":
        summary = discount_orders.aggregate(
            offer_discount=Coalesce(
                Sum("offer_discount"),
                Value(ZERO_MONEY),
                output_field=DecimalField(
                    max_digits=15,
                    decimal_places=2,
                ),
            ),
            coupon_discount=Coalesce(
                Sum("coupon_discount"),
                Value(ZERO_MONEY),
                output_field=DecimalField(
                    max_digits=15,
                    decimal_places=2,
                ),
            ),
            total_discount=Coalesce(
                Sum("discount_amount"),
                Value(ZERO_MONEY),
                output_field=DecimalField(
                    max_digits=15,
                    decimal_places=2,
                ),
            ),
        )

        # Keep total consistent even for older rows where discount_amount
        # was not populated correctly.
        calculated_total = (
            _money(summary["offer_discount"])
            + _money(summary["coupon_discount"])
        )

        stored_total = _money(summary["total_discount"])

        return {
            "offer_discount": _money(summary["offer_discount"]),
            "coupon_discount": _money(summary["coupon_discount"]),
            "total_discount": (
                stored_total
                if stored_total == calculated_total
                else calculated_total
            ),
        }

    try:
        selected_category_id = int(category_id)
    except (TypeError, ValueError):
        return {
            "offer_discount": ZERO_MONEY,
            "coupon_discount": ZERO_MONEY,
            "total_discount": ZERO_MONEY,
        }

    allocated_offer_discount = ZERO_MONEY
    allocated_coupon_discount = ZERO_MONEY

    for order in discount_orders:
        all_items = [
            item
            for item in order.items.all()
            if not item.is_cancelled
        ]

        order_item_value = sum(
            (_money(item.item_total) for item in all_items),
            ZERO_MONEY,
        )

        if order_item_value <= 0:
            continue

        category_item_value = sum(
            (
                _money(item.item_total)
                for item in all_items
                if (
                    item.variant
                    and item.variant.product
                    and item.variant.product.category_id
                    == selected_category_id
                )
            ),
            ZERO_MONEY,
        )

        if category_item_value <= 0:
            continue

        category_ratio = category_item_value / order_item_value

        allocated_offer_discount += _money(
            _money(order.offer_discount) * category_ratio
        )

        allocated_coupon_discount += _money(
            _money(order.coupon_discount) * category_ratio
        )

    allocated_offer_discount = _money(allocated_offer_discount)
    allocated_coupon_discount = _money(allocated_coupon_discount)

    return {
        "offer_discount": allocated_offer_discount,
        "coupon_discount": allocated_coupon_discount,
        "total_discount": _money(
            allocated_offer_discount + allocated_coupon_discount
        ),
    }


def get_sales_report_data(request):
    today = timezone.localdate()

    start_date_raw = request.GET.get("start_date", "").strip()
    end_date_raw = request.GET.get("end_date", "").strip()
    category_id = request.GET.get("category", "all").strip()
    search = request.GET.get("search", "").strip()
    sort = request.GET.get("sort", "revenue_high").strip()

    start_date = parse_report_date(start_date_raw)
    end_date = parse_report_date(end_date_raw)

    if not start_date:
        start_date = today.replace(day=1)

    if not end_date:
        end_date = today

    if start_date > end_date:
        start_date, end_date = end_date, start_date

        messages.warning(
            request,
            "Start date was after end date, so the dates were automatically corrected.",
        )

    report_days = (end_date - start_date).days + 1
    previous_end_date = start_date - timezone.timedelta(days=1)
    previous_start_date = (
        previous_end_date
        - timezone.timedelta(days=report_days - 1)
    )

    sales_items = (
        OrderItem.objects
        .filter(
            order__status__in=SALES_ORDER_STATUSES,
            order__ordered_at__date__range=[start_date, end_date],
            is_cancelled=False,
        )
        .select_related(
            "order",
            "order__user",
            "product",
            "variant",
            "variant__product",
            "variant__product__category",
        )
    )

    previous_items = (
        OrderItem.objects
        .filter(
            order__status__in=SALES_ORDER_STATUSES,
            order__ordered_at__date__range=[
                previous_start_date,
                previous_end_date,
            ],
            is_cancelled=False,
        )
        .select_related(
            "order",
            "order__user",
            "product",
            "variant",
            "variant__product",
            "variant__product__category",
        )
    )

    if category_id != "all":
        try:
            selected_category_id = int(category_id)

            sales_items = sales_items.filter(
                variant__product__category_id=selected_category_id
            )

            previous_items = previous_items.filter(
                variant__product__category_id=selected_category_id
            )
        except (TypeError, ValueError):
            category_id = "all"

    if search:
        search_filter = _build_search_filter(search)
        sales_items = sales_items.filter(search_filter)
        previous_items = previous_items.filter(search_filter)

    current_order_ids = sales_items.values_list(
        "order_id",
        flat=True,
    ).distinct()

    previous_order_ids = previous_items.values_list(
        "order_id",
        flat=True,
    ).distinct()

    sales_orders = (
        Order.objects
        .filter(
            id__in=current_order_ids,
            status__in=SALES_ORDER_STATUSES,
        )
        .select_related("user")
        .distinct()
    )

    previous_orders = (
        Order.objects
        .filter(
            id__in=previous_order_ids,
            status__in=SALES_ORDER_STATUSES,
        )
        .select_related("user")
        .distinct()
    )

    item_summary = sales_items.aggregate(
        total_revenue=Coalesce(
            Sum("item_total"),
            Value(ZERO_MONEY),
            output_field=DecimalField(
                max_digits=15,
                decimal_places=2,
            ),
        ),
        total_products_sold=Coalesce(
            Sum("quantity"),
            Value(0),
            output_field=IntegerField(),
        ),
    )

    previous_item_summary = previous_items.aggregate(
        total_revenue=Coalesce(
            Sum("item_total"),
            Value(ZERO_MONEY),
            output_field=DecimalField(
                max_digits=15,
                decimal_places=2,
            ),
        ),
        total_products_sold=Coalesce(
            Sum("quantity"),
            Value(0),
            output_field=IntegerField(),
        ),
    )

    total_revenue = _money(item_summary["total_revenue"])
    total_products_sold = item_summary["total_products_sold"] or 0

    previous_revenue = _money(
        previous_item_summary["total_revenue"]
    )
    previous_products_sold = (
        previous_item_summary["total_products_sold"] or 0
    )

    completed_orders = sales_orders.count()
    total_orders = completed_orders
    previous_order_count = previous_orders.count()

    discount_summary = _get_discount_summary(
        start_date=start_date,
        end_date=end_date,
        search=search,
        category_id=category_id,
    )

    offer_discount = discount_summary["offer_discount"]
    coupon_discount = discount_summary["coupon_discount"]
    total_discount = discount_summary["total_discount"]

    average_order_value = (
        total_revenue / completed_orders
        if completed_orders
        else ZERO_MONEY
    )

    previous_average_order_value = (
        previous_revenue / previous_order_count
        if previous_order_count
        else ZERO_MONEY
    )

    average_order_value = _money(average_order_value)
    previous_average_order_value = _money(
        previous_average_order_value
    )

    revenue_change = calculate_percentage_change(
        total_revenue,
        previous_revenue,
    )

    product_change = calculate_percentage_change(
        total_products_sold,
        previous_products_sold,
    )

    order_change = calculate_percentage_change(
        completed_orders,
        previous_order_count,
    )

    average_order_change = calculate_percentage_change(
        average_order_value,
        previous_average_order_value,
    )

    monthly_sales_queryset = (
        sales_items
        .annotate(
            month=TruncMonth("order__ordered_at")
        )
        .values("month")
        .annotate(
            revenue=Coalesce(
                Sum("item_total"),
                Value(ZERO_MONEY),
                output_field=DecimalField(
                    max_digits=15,
                    decimal_places=2,
                ),
            ),
            orders=Count(
                "order_id",
                distinct=True,
            ),
        )
        .order_by("month")
    )

    monthly_labels = []
    monthly_revenue = []
    monthly_orders = []

    for row in monthly_sales_queryset:
        if not row["month"]:
            continue

        monthly_labels.append(
            row["month"].strftime("%b %Y")
        )
        monthly_revenue.append(
            float(row["revenue"] or 0)
        )
        monthly_orders.append(
            row["orders"] or 0
        )

    category_sales_queryset = (
        sales_items
        .exclude(variant__isnull=True)
        .exclude(variant__product__isnull=True)
        .exclude(variant__product__category__isnull=True)
        .values(
            "variant__product__category_id",
            "variant__product__category__name",
        )
        .annotate(
            revenue=Coalesce(
                Sum("item_total"),
                Value(ZERO_MONEY),
                output_field=DecimalField(
                    max_digits=15,
                    decimal_places=2,
                ),
            ),
            quantity_sold=Coalesce(
                Sum("quantity"),
                Value(0),
                output_field=IntegerField(),
            ),
            completed_orders=Count(
                "order_id",
                distinct=True,
            ),
        )
        .order_by("-revenue")
    )

    category_labels = []
    category_revenue = []
    category_quantity = []

    for row in category_sales_queryset:
        category_labels.append(
            row["variant__product__category__name"]
            or "Uncategorized"
        )
        category_revenue.append(
            float(row["revenue"] or 0)
        )
        category_quantity.append(
            row["quantity_sold"] or 0
        )

    product_report = (
        sales_items
        .exclude(variant__isnull=True)
        .exclude(variant__product__isnull=True)
        .values(
            "variant__product_id",
            "variant__product__sku",
            "variant__product__name",
            "variant__product__category__name",
            "variant__product__total_stock",
        )
        .annotate(
            quantity_sold=Coalesce(
                Sum("quantity"),
                Value(0),
                output_field=IntegerField(),
            ),
            gross_revenue=Coalesce(
                Sum("item_total"),
                Value(ZERO_MONEY),
                output_field=DecimalField(
                    max_digits=15,
                    decimal_places=2,
                ),
            ),
            completed_order_count=Count(
                "order_id",
                distinct=True,
            ),
            last_sale_date=Max(
                "order__ordered_at"
            ),
        )
        .annotate(
            average_price=Coalesce(
                ExpressionWrapper(
                    F("gross_revenue")
                    / F("quantity_sold"),
                    output_field=DecimalField(
                        max_digits=12,
                        decimal_places=2,
                    ),
                ),
                Value(ZERO_MONEY),
                output_field=DecimalField(
                    max_digits=12,
                    decimal_places=2,
                ),
            )
        )
    )

    if sort == "revenue_low":
        product_report = product_report.order_by(
            "gross_revenue",
            "variant__product__name",
        )
    elif sort == "quantity_high":
        product_report = product_report.order_by(
            "-quantity_sold",
            "variant__product__name",
        )
    elif sort == "quantity_low":
        product_report = product_report.order_by(
            "quantity_sold",
            "variant__product__name",
        )
    elif sort == "name":
        product_report = product_report.order_by(
            "variant__product__name"
        )
    else:
        product_report = product_report.order_by(
            "-gross_revenue",
            "variant__product__name",
        )

    normalized_report = []

    for row in product_report:
        normalized_report.append({
            "report_product_id": row["variant__product_id"],
            "report_product_sku": row["variant__product__sku"],
            "report_product_name": row["variant__product__name"],
            "report_category_name": (
                row["variant__product__category__name"]
            ),
            "report_total_stock": (
                row["variant__product__total_stock"]
            ),
            "quantity_sold": row["quantity_sold"],
            "gross_revenue": row["gross_revenue"],
            "average_price": row["average_price"],
            "last_sale_date": row["last_sale_date"],
            "completed_order_count": (
                row["completed_order_count"]
            ),

            # Export compatibility keys
            "product_id": row["variant__product_id"],
            "product__sku": row["variant__product__sku"],
            "product__name": row["variant__product__name"],
            "product__category__name": (
                row["variant__product__category__name"]
            ),
            "product__total_stock": (
                row["variant__product__total_stock"]
            ),
        })

    categories = (
        Category.objects
        .filter(is_active=True)
        .order_by("name")
    )

    return {
        "sales_orders": sales_orders,
        "sales_items": sales_items,
        "product_report": normalized_report,
        "categories": categories,

        "start_date": start_date,
        "end_date": end_date,
        "category_id": category_id,
        "search": search,
        "sort": sort,

        "total_revenue": total_revenue,
        "total_orders": total_orders,
        "completed_orders": completed_orders,

        "offer_discount": offer_discount,
        "coupon_discount": coupon_discount,
        "total_discount": total_discount,

        "total_products_sold": total_products_sold,
        "average_order_value": average_order_value,

        "revenue_change": revenue_change,
        "product_change": product_change,
        "order_change": order_change,
        "average_order_change": average_order_change,

        "monthly_labels": monthly_labels,
        "monthly_revenue": monthly_revenue,
        "monthly_orders": monthly_orders,

        "category_labels": category_labels,
        "category_revenue": category_revenue,
        "category_quantity": category_quantity,
    }


@staff_member_required(login_url="admin_login")
def sales_report(request):
    report_data = get_sales_report_data(request)

    paginator = Paginator(
        report_data["product_report"],
        10,
    )

    page_obj = paginator.get_page(
        request.GET.get("page")
    )

    query_params = request.GET.copy()
    query_params.pop("page", None)
    query_params.pop("export", None)

    context = {
        **report_data,
        "page_obj": page_obj,
        "base_query": urlencode(
            query_params,
            doseq=True,
        ),
    }

    return render(
        request,
        "adminpanel/sales_report/sales_report.html",
        context,
    )


@staff_member_required(login_url="admin_login")
def export_sales_excel(request):
    report_data = get_sales_report_data(request)
    product_report = list(
        report_data["product_report"]
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "WheelVerse Sales Report"

    gold_fill = PatternFill(
        fill_type="solid",
        fgColor="F2CA50",
    )

    dark_fill = PatternFill(
        fill_type="solid",
        fgColor="17150F",
    )

    header_font = Font(
        bold=True,
        color="17150F",
    )

    white_font = Font(
        color="FFFFFF",
    )

    title_font = Font(
        bold=True,
        size=18,
        color="F2CA50",
    )

    worksheet.merge_cells("A1:I1")
    worksheet["A1"] = "WHEELVERSE SALES REPORT"
    worksheet["A1"].font = title_font
    worksheet["A1"].fill = dark_fill
    worksheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    worksheet.merge_cells("A2:I2")
    worksheet["A2"] = (
        f"Period: {report_data['start_date']:%d %b %Y}"
        f" to {report_data['end_date']:%d %b %Y}"
    )
    worksheet["A2"].font = white_font
    worksheet["A2"].fill = dark_fill
    worksheet["A2"].alignment = Alignment(
        horizontal="center"
    )

    summary_rows = [
        [
            "Total Revenue",
            float(report_data["total_revenue"]),
        ],
        [
            "Products Sold",
            report_data["total_products_sold"],
        ],
        [
            "Total Orders",
            report_data["total_orders"],
        ],
        [
            "Completed Orders",
            report_data["completed_orders"],
        ],
        [
            "Average Order Value",
            float(report_data["average_order_value"]),
        ],
        [
            "Offer Discount",
            float(report_data["offer_discount"]),
        ],
        [
            "Coupon Discount",
            float(report_data["coupon_discount"]),
        ],
        [
            "Total Discount",
            float(report_data["total_discount"]),
        ],
    ]

    current_row = 4

    for label, value in summary_rows:
        worksheet.cell(current_row, 1, label)
        worksheet.cell(current_row, 2, value)
        worksheet.cell(
            current_row,
            1,
        ).font = Font(bold=True)
        current_row += 1

    current_row += 1

    headers = [
        "Product ID",
        "SKU",
        "Product",
        "Category",
        "Average Price",
        "Quantity Sold",
        "Orders",
        "Revenue",
        "Last Sale",
    ]

    header_row = current_row

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )
        cell.fill = gold_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center"
        )

    for product in product_report:
        current_row += 1

        worksheet.append([
            product["product_id"] or "-",
            product["product__sku"] or "-",
            product["product__name"] or "Deleted Product",
            (
                product["product__category__name"]
                or "Uncategorized"
            ),
            float(product["average_price"] or 0),
            product["quantity_sold"] or 0,
            product["completed_order_count"] or 0,
            float(product["gross_revenue"] or 0),
            (
                product["last_sale_date"].strftime(
                    "%d %b %Y"
                )
                if product["last_sale_date"]
                else "-"
            ),
        ])

    column_widths = {
        1: 14,
        2: 18,
        3: 30,
        4: 20,
        5: 18,
        6: 16,
        7: 12,
        8: 18,
        9: 18,
    }

    for column_number, width in column_widths.items():
        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = width

    worksheet.freeze_panes = f"A{header_row + 1}"

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        min_col=5,
        max_col=8,
    ):
        row[0].number_format = '₹#,##0.00'
        row[3].number_format = '₹#,##0.00'

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    filename = (
        f"wheelverse_sales_"
        f"{report_data['start_date']:%Y%m%d}_"
        f"{report_data['end_date']:%Y%m%d}.xlsx"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )

    workbook.save(response)
    return response


@staff_member_required(login_url="admin_login")
def export_sales_pdf(request):
    report_data = get_sales_report_data(request)

    product_report = list(
        report_data["product_report"]
    )

    buffer = io.BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "WheelVerseTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#D4AF37"),
        alignment=TA_CENTER,
    )

    subtitle_style = ParagraphStyle(
        "WheelVerseSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        textColor=colors.HexColor("#444444"),
        alignment=TA_CENTER,
    )

    story = [
        Paragraph(
            "WHEELVERSE SALES REPORT",
            title_style,
        ),
        Paragraph(
            (
                f"Report period: "
                f"{report_data['start_date']:%d %B %Y} - "
                f"{report_data['end_date']:%d %B %Y}"
            ),
            subtitle_style,
        ),
        Spacer(1, 8 * mm),
    ]

    summary_data = [
        [
            "Revenue",
            "Products",
            "Orders",
            "Average",
            "Offer Disc.",
            "Coupon Disc.",
            "Total Disc.",
        ],
        [
            f"Rs. {report_data['total_revenue']:,.2f}",
            f"{report_data['total_products_sold']:,}",
            f"{report_data['total_orders']:,}",
            f"Rs. {report_data['average_order_value']:,.2f}",
            f"Rs. {report_data['offer_discount']:,.2f}",
            f"Rs. {report_data['coupon_discount']:,.2f}",
            f"Rs. {report_data['total_discount']:,.2f}",
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[37 * mm] * 7,
    )

    summary_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#17150F"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.HexColor("#F2CA50"),
            ),
            (
                "BACKGROUND",
                (0, 1),
                (-1, 1),
                colors.HexColor("#F5F2E9"),
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER",
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.4,
                colors.HexColor("#B99B3D"),
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
        ])
    )

    story.extend([
        summary_table,
        Spacer(1, 8 * mm),
    ])

    report_table_data = [[
        "ID",
        "SKU",
        "Product",
        "Category",
        "Avg. Price",
        "Qty",
        "Orders",
        "Revenue",
        "Stock",
        "Last Sale",
    ]]

    for product in product_report:
        report_table_data.append([
            str(product["product_id"] or "-"),
            product["product__sku"] or "-",
            product["product__name"] or "Deleted Product",
            (
                product["product__category__name"]
                or "Uncategorized"
            ),
            f"Rs. {product['average_price'] or 0:,.2f}",
            str(product["quantity_sold"] or 0),
            str(product["completed_order_count"] or 0),
            f"Rs. {product['gross_revenue'] or 0:,.2f}",
            str(product["product__total_stock"] or 0),
            (
                product["last_sale_date"].strftime(
                    "%d %b %Y"
                )
                if product["last_sale_date"]
                else "-"
            ),
        ])

    report_table = Table(
        report_table_data,
        repeatRows=1,
        colWidths=[
            14 * mm,
            22 * mm,
            42 * mm,
            28 * mm,
            27 * mm,
            16 * mm,
            17 * mm,
            29 * mm,
            17 * mm,
            25 * mm,
        ],
    )

    report_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#17150F"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.HexColor("#F2CA50"),
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "FONTNAME",
                (0, 1),
                (-1, -1),
                "Helvetica",
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                8,
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.35,
                colors.HexColor("#B6B0A0"),
            ),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [
                    colors.white,
                    colors.HexColor("#F8F6EF"),
                ],
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE",
            ),
            (
                "ALIGN",
                (0, 0),
                (1, -1),
                "CENTER",
            ),
            (
                "ALIGN",
                (4, 1),
                (9, -1),
                "RIGHT",
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                6,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                6,
            ),
        ])
    )

    story.append(report_table)

    document.build(story)

    pdf_value = buffer.getvalue()
    buffer.close()

    response = HttpResponse(
        pdf_value,
        content_type="application/pdf",
    )

    filename = (
        f"wheelverse_sales_"
        f"{report_data['start_date']:%Y%m%d}_"
        f"{report_data['end_date']:%Y%m%d}.pdf"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )

    return response